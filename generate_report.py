#!/usr/bin/env python3
"""
Script tạo file báo cáo thực tập (.docx) với format chuẩn.
Cách dùng:
  1. pip install python-docx
  2. python generate_report.py
  3. Mở file bao_cao_thuc_tap.docx bằng Word
  4. Nhấn Ctrl+A → F9 để cập nhật mục lục
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import datetime


# ============================================================
# THÔNG TIN CÁ NHÂN - THAY ĐỔI TẠI ĐÂY
# ============================================================
INFO = {
    "truong": "TRƯỜNG ĐẠI HỌC ...",
    "khoa": "KHOA CÔNG NGHỆ THÔNG TIN",
    "de_tai": "XÂY DỰNG HỆ THỐNG CHUYỂN ĐỔI ĐỊA CHỈ\nVN ADDRESS CONVERTER",
    "ho_ten": "Nguyễn Văn A",
    "mssv": "20xxxxxx",
    "lop": "CNTT - K20",
    "gvhd": "ThS. Nguyễn Văn B",
    "nguoi_huong_dan_cty": "Anh/Chị C – Senior Developer",
    "cong_ty": "CÔNG TY TNHH GIẢI PHÁP THÔNG MINH POPIPLUS",
    "nam": "2026",
}


def set_cell_shading(cell, color_hex):
    """Tô màu nền cho ô trong bảng."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_paragraph_format(paragraph, font_name="Times New Roman", font_size=13,
                         bold=False, italic=False, alignment=None,
                         space_before=6, space_after=6, line_spacing=1.5,
                         first_line_indent=None, color=None):
    """Thiết lập format cho một đoạn văn."""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if alignment is not None:
        pf.alignment = alignment
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)

    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = RGBColor(*color)
        # Đảm bảo font hiển thị đúng trên Word
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_heading_custom(doc, text, level=1):
    """Thêm heading với format chuẩn báo cáo."""
    heading = doc.add_heading(text, level=level)

    sizes = {0: 16, 1: 15, 2: 14, 3: 13}
    font_size = sizes.get(level, 13)

    for run in heading.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)  # Màu đen
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")

    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.line_spacing = 1.5

    return heading


def add_paragraph_text(doc, text, bold=False, italic=False, indent=True,
                       alignment=None, font_size=13):
    """Thêm một đoạn văn bản với format chuẩn."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_paragraph_format(
        p,
        font_size=font_size,
        bold=bold,
        italic=italic,
        alignment=alignment,
        first_line_indent=1.27 if indent else None,
    )
    return p


def add_bullet(doc, text, level=0):
    """Thêm một mục gạch đầu dòng."""
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    if level > 0:
        p.paragraph_format.left_indent = Cm(1.27 * (level + 1))
    return p


def add_table_standard(doc, headers, rows, caption=None):
    """
    Thêm bảng chuẩn báo cáo.
    - caption: "Bảng X.Y: Mô tả"
    - headers: danh sách tên cột
    - rows: danh sách các dòng (mỗi dòng là list)
    """
    # Tiêu đề bảng (phía trên)
    if caption:
        cap_p = doc.add_paragraph()
        # Phần "Bảng X.Y:" in đậm
        parts = caption.split(":", 1)
        run1 = cap_p.add_run(parts[0] + ":")
        run1.font.name = "Times New Roman"
        run1.font.size = Pt(12)
        run1.font.bold = True
        run1._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        if len(parts) > 1:
            run2 = cap_p.add_run(parts[1])
            run2.font.name = "Times New Roman"
            run2.font.size = Pt(12)
            run2.font.italic = True
            run2._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_after = Pt(4)

    # Tạo bảng
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        set_cell_shading(cell, "D9E2F3")

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

    # Nguồn
    src_p = doc.add_paragraph()
    run = src_p.add_run("Nguồn: Tác giả tự tổng hợp")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.italic = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
    src_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    src_p.paragraph_format.space_before = Pt(2)
    src_p.paragraph_format.space_after = Pt(10)

    return table


def add_toc(doc):
    """Thêm mục lục tự động (cập nhật bằng Ctrl+A → F9 trong Word)."""
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fld_char_begin)

    run2 = p.add_run()
    instr_text = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._element.append(instr_text)

    run3 = p.add_run()
    fld_char_separate = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._element.append(fld_char_separate)

    run4 = p.add_run("(Nhấn Ctrl+A → F9 để cập nhật mục lục)")
    run4.font.name = "Times New Roman"
    run4.font.size = Pt(13)
    run4.font.color.rgb = RGBColor(128, 128, 128)

    run5 = p.add_run()
    fld_char_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._element.append(fld_char_end)


def add_page_number(doc):
    """Thêm số trang ở giữa, phía dưới."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run()
        fld_char_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        run._element.append(fld_char_begin)

        run2 = p.add_run()
        instr_text = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        run2._element.append(instr_text)

        run3 = p.add_run()
        fld_char_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        run3._element.append(fld_char_end)


def create_cover_page(doc):
    """Tạo trang bìa."""
    # Tên trường
    p = doc.add_paragraph()
    run = p.add_run(INFO["truong"])
    set_paragraph_format(p, font_size=14, bold=True,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=0, space_after=0)

    # Tên khoa
    p = doc.add_paragraph()
    run = p.add_run(INFO["khoa"])
    set_paragraph_format(p, font_size=14, bold=True,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=0, space_after=36)

    # Khoảng trống
    for _ in range(3):
        doc.add_paragraph()

    # Tiêu đề "BÁO CÁO THỰC TẬP"
    p = doc.add_paragraph()
    run = p.add_run("BÁO CÁO THỰC TẬP TỐT NGHIỆP")
    set_paragraph_format(p, font_size=20, bold=True,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=0, space_after=12)

    # Tên đề tài
    p = doc.add_paragraph()
    run = p.add_run(INFO["de_tai"])
    set_paragraph_format(p, font_size=16, bold=True,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         space_before=12, space_after=36)

    # Khoảng trống
    for _ in range(3):
        doc.add_paragraph()

    # Thông tin sinh viên
    info_lines = [
        f"Họ và tên:  {INFO['ho_ten']}",
        f"MSSV:  {INFO['mssv']}",
        f"Lớp:  {INFO['lop']}",
        f"GVHD:  {INFO['gvhd']}",
        f"Đơn vị thực tập:  {INFO['cong_ty']}",
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        # Split at first ":"
        parts = line.split(":", 1)
        r1 = p.add_run(parts[0] + ":")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        r2 = p.add_run(parts[1])
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(13)
        r2._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.left_indent = Cm(4)

    # Khoảng trống
    for _ in range(2):
        doc.add_paragraph()

    # Năm
    p = doc.add_paragraph()
    run = p.add_run(f"TP. Hồ Chí Minh, {INFO['nam']}")
    set_paragraph_format(p, font_size=13, bold=False, italic=True,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()


def create_acknowledgment(doc):
    """Tạo trang Lời cảm ơn."""
    add_heading_custom(doc, "LỜI CẢM ƠN", level=0)

    add_paragraph_text(doc,
        "Trong quá trình thực tập tại Công ty TNHH Giải Pháp Thông Minh POPIPLUS, "
        "em đã nhận được sự hướng dẫn tận tình từ quý thầy cô và các anh chị "
        "trong công ty. Em xin chân thành cảm ơn:"
    )
    add_bullet(doc, f"{INFO['gvhd']} – đã hướng dẫn, góp ý và theo dõi tiến độ thực tập.")
    add_bullet(doc, f"{INFO['nguoi_huong_dan_cty']} – đã trực tiếp hướng dẫn kỹ thuật, "
                    "review code và chia sẻ kinh nghiệm làm việc thực tế.")
    add_bullet(doc, "Ban lãnh đạo và các thành viên công ty POPIPLUS – đã tạo điều kiện "
                    "cho em tham gia dự án thực tế.")
    add_bullet(doc, "Gia đình và bạn bè – đã luôn động viên và hỗ trợ em trong suốt "
                    "quá trình học tập và thực tập.")

    add_paragraph_text(doc,
        "Em xin gửi lời cảm ơn sâu sắc nhất đến tất cả mọi người. "
        "Do kiến thức và kinh nghiệm còn hạn chế, báo cáo chắc chắn không tránh khỏi "
        "những thiếu sót. Em rất mong nhận được sự góp ý từ quý thầy cô để hoàn thiện hơn."
    )

    p = doc.add_paragraph()
    run = p.add_run(f"\nTP. Hồ Chí Minh, ngày ... tháng ... năm {INFO['nam']}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.italic = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p = doc.add_paragraph()
    run = p.add_run(INFO["ho_ten"])
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()


def create_review_pages(doc):
    """Tạo trang nhận xét của GVHD và công ty."""
    for title in ["NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN", "NHẬN XÉT CỦA ĐƠN VỊ THỰC TẬP"]:
        add_heading_custom(doc, title, level=0)
        # Để trống cho người viết nhận xét
        for _ in range(20):
            p = doc.add_paragraph()
            run = p.add_run("." * 90)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(220, 220, 220)
            p.paragraph_format.line_spacing = 2.0

        p = doc.add_paragraph()
        run = p.add_run(f"TP. Hồ Chí Minh, ngày ... tháng ... năm {INFO['nam']}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.italic = True
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "Times New Roman")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_page_break()


def create_toc_page(doc):
    """Tạo trang mục lục."""
    add_heading_custom(doc, "MỤC LỤC", level=0)
    add_toc(doc)
    doc.add_page_break()


# ============================================================
# CHƯƠNG 1
# ============================================================
def create_chapter1(doc):
    add_heading_custom(doc, "CHƯƠNG 1: TỔNG QUAN VỀ ĐƠN VỊ THỰC TẬP", level=1)

    # 1.1
    add_heading_custom(doc, "1.1. Giới thiệu đơn vị thực tập", level=2)
    add_paragraph_text(doc,
        "CÔNG TY TNHH GIẢI PHÁP THÔNG MINH POPIPLUS (Mã số doanh nghiệp: 0318272823) "
        "được cấp giấy chứng nhận đăng ký kinh doanh ngày 19/01/2024 bởi Sở Kế hoạch "
        "và Đầu tư Thành phố Hồ Chí Minh. Trụ sở đặt tại Thành phố Hồ Chí Minh."
    )
    add_paragraph_text(doc,
        "Nền tảng chủ lực của công ty là LOZIDO – nền tảng số hoá quản lý nhà trọ, "
        "phòng trọ và việc làm. POPIPLUS được thành lập với mục tiêu xây dựng một hệ "
        "sinh thái số hoá toàn diện, giúp người dùng nhanh chóng tìm kiếm, đăng tin "
        "và quản lý bất động sản cũng như việc làm một cách an toàn và hiệu quả."
    )
    add_paragraph_text(doc,
        "Công ty hiện có quy mô khoảng 15–20 nhân viên, chia thành 3–4 nhóm phát triển "
        "song song. Mỗi nhóm gồm từ 3 đến 5 thành viên, bao gồm trưởng nhóm, backend "
        "developer, frontend developer và tester. Dự án VN Address Converter được giao "
        "cho một nhóm gồm 3 thành viên: 2 backend developer và 1 frontend developer, "
        "dưới sự hướng dẫn trực tiếp của mentor (Senior Developer)."
    )

    # 1.2
    add_heading_custom(doc, "1.2. Lĩnh vực hoạt động", level=2)
    add_table_standard(doc,
        headers=["Lĩnh vực", "Mô tả ngắn gọn"],
        rows=[
            ["Nền tảng tìm trọ & căn hộ",
             "Phát triển website và ứng dụng di động cho phép người dùng đăng tin, "
             "tìm kiếm và liên hệ trực tiếp với chủ nhà hoặc môi giới."],
            ["Dịch vụ tuyển dụng",
             "Cung cấp công cụ đăng tuyển, tìm việc và quản lý hồ sơ ứng viên."],
            ["Công nghệ dữ liệu địa lý",
             "Xây dựng cơ sở dữ liệu địa chỉ, tích hợp bản đồ và các API tra cứu."],
            ["Giải pháp phần mềm tùy chỉnh",
             "Thiết kế, phát triển và bảo trì các hệ thống phần mềm theo yêu cầu."],
        ],
        caption="Bảng 1.1: Các lĩnh vực hoạt động của POPIPLUS"
    )

    # 1.3
    add_heading_custom(doc, "1.3. Văn hóa và môi trường làm việc", level=2)
    add_paragraph_text(doc,
        "POPIPLUS hoạt động theo mô hình làm việc từ xa (remote), với các giá trị "
        "cốt lõi: Sáng tạo – Chất lượng – Trách nhiệm – Hợp tác."
    )
    add_bullet(doc, "Kênh giao tiếp chính: Slack (chat nhanh, tạo kênh dự án, chia sẻ tài liệu) "
                    "và Google Meet (họp hằng ngày, review code, demo tính năng).")
    add_bullet(doc, "Quản lý công việc: Trello – board Kanban để tạo, phân công và theo dõi task.")
    add_bullet(doc, "Môi trường phát triển đồng nhất: Docker hoặc VS Code Remote Containers.")
    add_bullet(doc, "Giờ làm việc linh hoạt: Thống nhất khung giờ 9h–12h, 14h–18h.")
    add_bullet(doc, "Đào tạo: Workshop/webinar hàng tháng, mentoring trực tuyến 1:1.")
    add_bullet(doc, "Hoạt động gắn kết: Virtual coffee break mỗi tuần, online hackathon mỗi 6 tháng.")
    add_bullet(doc, "Bảo mật: VPN/SSH key, MFA cho tài khoản Git, cloud và công cụ quản lý.")

    # 1.4
    add_heading_custom(doc, "1.4. Các quy trình làm việc", level=2)

    add_heading_custom(doc, "1.4.1. Đề xuất dự án", level=3)
    add_paragraph_text(doc,
        "Thành viên đưa ra ý tưởng dự án hoặc tính năng mới trên Trello (thẻ \"Idea\"). "
        "Trưởng nhóm xem xét, bổ sung thông tin chi tiết (mục tiêu, phạm vi, lợi ích). "
        "Tính khả thi được đánh giá trong buổi Sprint Planning trên Google Meet. "
        "Khi được phê duyệt, thẻ chuyển sang cột \"Backlog\" và gán nhãn Priority tương ứng."
    )

    add_heading_custom(doc, "1.4.2. Phân tích yêu cầu", level=3)
    add_paragraph_text(doc,
        "Thu thập yêu cầu từ khách hàng hoặc bộ phận kinh doanh qua Slack và Google Docs. "
        "Viết User Story trong Trello (cột \"To Do\") kèm Acceptance Criteria rõ ràng. "
        "Buổi họp ngắn (15–30 phút) trên Meet để xác nhận yêu cầu, "
        "sau đó lưu trữ tài liệu trên Confluence."
    )

    add_heading_custom(doc, "1.4.3. Thiết kế kiến trúc", level=3)
    add_paragraph_text(doc,
        "Sử dụng Excalidraw để vẽ sơ đồ quan hệ (ERD), lưu trong thư mục design/ trên repo. "
        "API Specification được tạo bằng OpenAPI/Swagger, chia sẻ qua GitHub và Slack. "
        "UI/UX Wireframe thiết kế nhanh trong Figma, liên kết vào thẻ Trello để đội frontend tham khảo."
    )

    add_heading_custom(doc, "1.4.4. Phát triển (Sprint)", level=3)
    add_paragraph_text(doc,
        "Mỗi sprint kéo dài 1 tuần, bắt đầu bằng Sprint Planning (Google Meet) "
        "và kết thúc bằng Sprint Review + Retrospective. "
        "Mỗi tính năng có một branch feature/<tên tính năng>; tạo Pull Request ngay khi hoàn thành. "
        "Các thành viên khác review PR trên GitHub, để lại comment trong Slack. "
        "Khi PR được merge, GitHub Actions tự động chạy unit test, lint và deploy lên Staging."
    )

    add_heading_custom(doc, "1.4.5. Kiểm thử", level=3)
    add_paragraph_text(doc,
        "Unit Test viết bằng Jest (backend) và React Testing Library (frontend). "
        "Integration Test sử dụng Supertest cho API, chạy trong pipeline CI. "
        "E2E Test thực hiện bằng Cypress, lưu kết quả trong thư mục cypress/reports. "
        "Khi phát hiện lỗi, tạo thẻ \"Bug\" trên Trello, gán người chịu trách nhiệm và đặt deadline."
    )

    add_heading_custom(doc, "1.4.6. Triển khai (Release)", level=3)
    add_paragraph_text(doc,
        "Sau khi CI passes, hệ thống tự động deploy lên server staging (Docker). "
        "Nhóm QA thực hiện kiểm tra nhanh (Smoke Test). "
        "Khi QA ký duyệt, thực hiện Release bằng GitHub Actions tới môi trường production. "
        "Giám sát log và metric bằng Grafana/Prometheus, thông báo bất thường qua Slack."
    )

    add_heading_custom(doc, "1.4.7. Bảo trì & Hỗ trợ", level=3)
    add_paragraph_text(doc,
        "Khi có sự cố, mở thẻ \"Incident\" trên Trello, ghi lại thời gian, nguyên nhân và bước "
        "khắc phục. Các bản vá được đưa vào branch hotfix/<tên vấn đề>, tạo PR và triển khai nhanh. "
        "Mọi thay đổi đều phải cập nhật tài liệu trên Confluence và API docs."
    )

    add_heading_custom(doc, "1.4.8. Đánh giá & Cải tiến (Retrospective)", level=3)
    add_paragraph_text(doc,
        "Buổi Retrospective diễn ra vào cuối mỗi sprint trên Google Meet, mọi thành viên "
        "chia sẻ \"What went well\", \"What didn't go well\" và \"Action items\". "
        "Các hành động cải tiến được ghi lại trong Trello (cột \"Improvement\") và theo dõi tiến độ. "
        "Áp dụng các đề xuất vào quy trình kế tiếp, cập nhật SOP trong Confluence."
    )

    # 1.5
    add_heading_custom(doc, "1.5. Vai trò của sinh viên thực tập", level=2)
    add_paragraph_text(doc,
        f"Trong thời gian thực tập, tôi đảm nhận vị trí Backend Developer trong nhóm phát triển "
        f"dự án VN Address Converter. Nhiệm vụ chính của tôi là thiết kế và xây dựng các API "
        f"phục vụ cho việc chuyển đổi địa chỉ, tìm kiếm mờ và gợi ý tự động. Cụ thể, tôi chịu "
        f"trách nhiệm phát triển 5 service cốt lõi (dropdown, fuzzy search, newToOld mapping, "
        f"oldToNew mapping, suggest) và hệ thống xác thực người dùng (authentication). "
        f"Ngoài ra, tôi cũng tham gia viết unit test, integration test và cập nhật tài liệu "
        f"Swagger cho toàn bộ API. Mentor hướng dẫn là {INFO['nguoi_huong_dan_cty']}, "
        f"hỗ trợ review code và giải đáp các vấn đề kỹ thuật qua các buổi meeting 1:1 "
        f"hằng tuần trên Google Meet."
    )

    doc.add_page_break()


# ============================================================
# CHƯƠNG 2
# ============================================================
def create_chapter2(doc):
    add_heading_custom(doc, "CHƯƠNG 2: BÀI TOÁN / ĐỀ TÀI / DỰ ÁN", level=1)

    # 2.1
    add_heading_custom(doc, "2.1. Thực trạng và vấn đề hiện tại đang tồn đọng", level=2)
    add_paragraph_text(doc,
        "Trong dự án này, tôi đảm nhận vai trò Backend Developer. Dưới đây là phân tích "
        "thực trạng hệ thống, trong đó phần backend do tôi trực tiếp khảo sát và đánh giá, "
        "phần frontend được ghi nhận tổng quan để có cái nhìn toàn diện về dự án."
    )
    add_paragraph_text(doc,
        "Hiện tại hệ thống backend vẫn còn một số hạn chế đáng chú ý. Đầu tiên, việc kết nối "
        "cơ sở dữ liệu được thực hiện bằng pg.Client thay vì một pool kết nối, dẫn đến chỉ có "
        "một kết nối duy nhất và gây nghẽn khi có nhiều yêu cầu đồng thời. Thêm vào đó, các API "
        "trả về toàn bộ dữ liệu mà không có cơ chế phân trang, khiến phản hồi trở nên nặng và "
        "chậm khi dữ liệu tăng lên. Việc ghi log hiện đang chỉ dùng console.log đơn giản, "
        "không có mức độ log, không ghi vào file và không có timestamp."
    )
    add_paragraph_text(doc,
        "Hệ thống chưa có các unit test cho service và model, vì vậy khi thực hiện refactor "
        "sẽ không có cơ chế kiểm tra tự động. Một số hàm xác thực và các service ánh xạ "
        "cũ–mới gần như trùng lặp, gây lãng phí công sức bảo trì. Truy vấn fuzzy search còn "
        "quá phức tạp với tới tám JOIN, khó debug và chưa tối ưu hiệu năng. API công khai "
        "chưa có cơ chế giới hạn tần suất (rate limiting), dễ bị lạm dụng hoặc tấn công DDoS. "
        "Cấu trúc bảng dữ liệu được tạo bằng file .sql thủ công, không có công cụ migration."
    )
    add_paragraph_text(doc,
        "Về phía frontend (do thành viên khác trong nhóm phát triển), trang quản trị hiện vẫn "
        "sử dụng dữ liệu giả (mock data) và chưa gọi API thực tế. Các nút Thêm, Sửa, Xóa chưa "
        "có handler. Khi token hết hạn, ứng dụng chưa có cơ chế tự động logout. Trang Profile "
        "chưa được xây dựng. Chức năng hiển thị bản đồ còn gặp vấn đề về độ chính xác khi "
        "địa chỉ ở cấp xã/phường."
    )

    # 2.2
    add_heading_custom(doc, "2.2. Phát biểu bài toán / đề tài / dự án", level=2)
    add_paragraph_text(doc,
        "Mục tiêu của đề tài là xây dựng hệ thống VN Address Converter bao gồm backend "
        "và frontend, đáp ứng các yêu cầu sau:"
    )
    add_bullet(doc, "Cho phép người dùng tra cứu, gợi ý và tìm kiếm các đơn vị hành chính "
                    "(tỉnh, huyện, xã) hiện hành và lịch sử thay đổi.")
    add_bullet(doc, "Hỗ trợ chuyển đổi tự động giữa địa chỉ cũ và địa chỉ mới dựa trên "
                    "bảng ánh xạ thay đổi do nhà nước công bố.")
    add_bullet(doc, "Cung cấp một API chuẩn OpenAPI/Swagger để các hệ thống bên thứ ba "
                    "(cơ quan nhà nước, doanh nghiệp) có thể tích hợp nhanh chóng.")
    add_bullet(doc, "Đảm bảo bảo mật toàn diện bằng JWT, phân quyền (RBAC), HTTPS và "
                    "giới hạn tần suất truy cập.")
    add_bullet(doc, "Giám sát hoạt động hệ thống bằng Grafana và Prometheus, đồng thời "
                    "thông báo kết quả triển khai qua Slack.")

    add_paragraph_text(doc,
        "Đối tượng sử dụng hệ thống bao gồm: cơ quan nhà nước cần chuyển đổi địa chỉ "
        "theo nghị quyết mới, doanh nghiệp logistics và bất động sản cần chuẩn hoá dữ liệu "
        "địa chỉ, và quản trị viên (admin) chịu trách nhiệm quản lý dữ liệu địa lý trong hệ thống."
    )

    # 2.3
    add_heading_custom(doc, "2.3. Đề xuất các giải pháp kiến nghị để giải quyết bài toán", level=2)
    add_paragraph_text(doc,
        "Để khắc phục các vấn đề đã nêu và đạt được mục tiêu đề ra, chúng tôi đề xuất "
        "các giải pháp sau:"
    )
    add_bullet(doc, "Cải thiện kết nối cơ sở dữ liệu: Thay pg.Client bằng pg.Pool hoặc "
                    "sử dụng Prisma để quản lý pool kết nối, giúp hệ thống chịu tải tốt hơn.")
    add_bullet(doc, "Thêm phân trang cho API: Đưa các tham số page và limit vào các endpoint "
                    "danh sách, giảm tải mạng và cải thiện thời gian phản hồi.")
    add_bullet(doc, "Triển khai hệ thống log chuẩn: Sử dụng thư viện Winston với các mức "
                    "log (info, warn, error) và ghi log vào file, bổ sung timestamp.")
    add_bullet(doc, "Xây dựng unit test: Viết các test cho service, controller và model bằng "
                    "Jest và Supertest, đạt coverage ít nhất 80%.")
    add_bullet(doc, "Tái cấu trúc mã lặp: Gom các hàm xác thực giống nhau thành một hàm chung, "
                    "tạo service chung cho việc ánh xạ cũ–mới và mới–cũ.")
    add_bullet(doc, "Tối ưu fuzzy search: Tạo view hoặc stored procedure trong PostgreSQL để "
                    "tách logic truy vấn, sử dụng GIN index cho pg_trgm.")
    add_bullet(doc, "Áp dụng rate limiting: Sử dụng express-rate-limit để giới hạn số request "
                    "theo IP ở tầng backend, bảo vệ hệ thống khỏi lạm dụng.")
    add_bullet(doc, "Quản lý migration tự động: Áp dụng Prisma Migrate hoặc Flyway để quản lý "
                    "phiên bản schema, tránh việc chạy file .sql thủ công.")

    # 2.4
    add_heading_custom(doc, "2.4. Xây dựng kế hoạch công việc cụ thể chi tiết để thực hiện", level=2)
    add_paragraph_text(doc,
        "Trong kế hoạch dưới đây, phần Backend là công việc do tôi trực tiếp thực hiện, "
        "bao gồm xây dựng API, viết test và cấu hình bảo mật. Phần Frontend do thành viên "
        "khác trong nhóm đảm nhận và được ghi nhận tổng quan để thể hiện sự phối hợp giữa "
        "backend và frontend trong mỗi sprint."
    )

    # Sprint 1
    add_heading_custom(doc, "Sprint 1 – Tuần 1: Service địa lý + Giao diện tương tác", level=3)
    add_table_standard(doc,
        headers=["Ngày", "Backend (tôi thực hiện)", "Frontend (TV khác)", "Kết quả mong đợi"],
        rows=[
            ["Thứ 2", "Thiết lập môi trường test (Jest, Supertest, Redis). Tạo nhánh sprint1.",
             "Thiết kế mock UI Dropdown (React + Tailwind).", "Môi trường chuẩn, UI mẫu sẵn sàng."],
            ["Thứ 3", "Refactor dropdown.service.js – validation, pagination, cache Redis, logging.",
             "Kết nối UI Dropdown với API.", "Dropdown hoạt động."],
            ["Thứ 4", "Refactor fuzzy.service.js – view/stored procedure, validation, pagination, cache.",
             "Thêm Search Bar, hiển thị kết quả fuzzy.", "Fuzzy search hoạt động."],
            ["Thứ 5", "Refactor newToOld.service.js & oldToNew.service.js – validation, pagination, cache.",
             "Thêm Mapping UI: 2 tab Mới→Cũ và Cũ→Mới.", "Mapping từ UI hoạt động."],
            ["Thứ 6", "Refactor suggest.service.js – validation, limit 10, cache, logging.",
             "Thêm Autocomplete cho ô nhập địa chỉ.", "Suggest hoạt động mượt."],
        ],
        caption="Bảng 2.1: Kế hoạch công việc Sprint 1"
    )

    # Sprint 2
    add_heading_custom(doc, "Sprint 2 – Tuần 2: Authentication & Bảo mật", level=3)
    add_table_standard(doc,
        headers=["Ngày", "Backend (tôi thực hiện)", "Frontend (TV khác)", "Kết quả mong đợi"],
        rows=[
            ["Thứ 2", "Thiết kế schema users, refresh_tokens (PostgreSQL). Tạo migration.",
             "Thiết kế Login / Register page.", "DB schema sẵn sàng, UI login cơ bản."],
            ["Thứ 3", "Implement auth.service.js (register, login, refresh, logout).",
             "Kết nối UI login/register với API.", "Đăng ký/đăng nhập hoạt động."],
            ["Thứ 4", "Thêm middleware JWT, rate-limit (express-rate-limit), Helmet, CORS.",
             "Thêm Route Guard bảo vệ trang admin.", "Bảo mật API và UI."],
            ["Thứ 5", "Viết unit test cho Auth Service (Jest).",
             "Viết e2e test (Cypress) cho luồng login.", "Test backend và UI."],
            ["Thứ 6", "Viết integration test (Supertest) cho toàn bộ luồng auth.",
             "Kiểm tra UI khi login thất bại, token hết hạn.", "Luồng auth ổn định."],
        ],
        caption="Bảng 2.2: Kế hoạch công việc Sprint 2"
    )

    # Liên kết Backend ↔ Frontend
    add_paragraph_text(doc,
        "Bảng dưới đây mô tả cách các API backend (do tôi xây dựng) được frontend "
        "(do thành viên khác phát triển) sử dụng, nhằm thể hiện sự liên kết giữa hai "
        "phần trong hệ thống.",
        indent=True
    )
    add_table_standard(doc,
        headers=["Backend Service", "Endpoint", "Frontend Component", "Mô tả tích hợp"],
        rows=[
            ["Dropdown", "GET /api/v1/dropdowns/provinces", "ProvinceDropdown",
             "Component gọi API, nhận dữ liệu pagination."],
            ["Fuzzy Search", "GET /api/v1/search?keyword=…", "FuzzySearchBar",
             "Nhập từ khóa → gọi API → hiển thị kết quả."],
            ["Mapping", "GET /api/v1/mapping/new-to-old", "MappingPanel",
             "Form nhập địa chỉ → gọi API → bảng mapping."],
            ["Suggest", "GET /api/v1/suggest?q=…", "AddressAutocomplete",
             "Gợi ý tự động khi nhập địa chỉ."],
            ["Auth", "/api/v1/auth/*", "LoginPage, ProtectedRoute",
             "JWT httpOnly cookie, route guard."],
        ],
        caption="Bảng 2.3: Liên kết Backend ↔ Frontend"
    )

    # Kết luận
    add_paragraph_text(doc,
        "Sprint 1 cung cấp đầy đủ các service địa lý và giao diện tương tác, đồng thời "
        "đã chuẩn hoá (validation, pagination, cache, logging, test, Swagger). Sprint 2 bổ "
        "sung hệ thống Authentication và bảo mật cho cả backend và frontend, cho phép người "
        "dùng đăng nhập và truy cập các chức năng đã xây dựng trong Sprint 1. Không có công "
        "việc nào được lên lịch vào thứ bảy và chủ nhật, mọi nhiệm vụ đều nằm trong các "
        "ngày làm việc từ thứ hai đến thứ sáu."
    )

    doc.add_page_break()


# ============================================================
# CHƯƠNG 3 (khung sẵn)
# ============================================================
def create_chapter3(doc):
    add_heading_custom(doc, "CHƯƠNG 3: PHÂN TÍCH, THIẾT KẾ VÀ TRIỂN KHAI THỰC HIỆN", level=1)

    add_heading_custom(doc, "3.1. Tóm lược các nội dung quan trọng cần thực hiện", level=2)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)

    add_heading_custom(doc, "3.2. Phân tích và thiết kế", level=2)
    add_heading_custom(doc, "3.2.1. Phân tích nghiệp vụ", level=3)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)
    add_heading_custom(doc, "3.2.2. Kiến trúc hệ thống (Node.js + Express)", level=3)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)
    add_heading_custom(doc, "3.2.3. Cơ sở dữ liệu", level=3)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)

    add_heading_custom(doc, "3.3. Đề xuất thực hiện", level=2)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)

    add_heading_custom(doc, "3.4. Thực hiện và đánh giá kết quả đạt được", level=2)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)

    doc.add_page_break()


# ============================================================
# CHƯƠNG 4, KẾT LUẬN, TÀI LIỆU THAM KHẢO (khung)
# ============================================================
def create_chapter4(doc):
    add_heading_custom(doc, "CHƯƠNG 4: KẾT QUẢ VÀ ĐÁNH GIÁ", level=1)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)
    doc.add_page_break()


def create_conclusion(doc):
    add_heading_custom(doc, "KẾT LUẬN VÀ KIẾN NGHỊ", level=1)
    add_paragraph_text(doc, "[Viết nội dung tại đây]", italic=True)
    doc.add_page_break()


def create_references(doc):
    add_heading_custom(doc, "TÀI LIỆU THAM KHẢO", level=1)
    refs = [
        "[1] Express.js Documentation, https://expressjs.com/",
        "[2] PostgreSQL Documentation, https://www.postgresql.org/docs/",
        "[3] Redis Documentation, https://redis.io/docs/",
        "[4] JSON Web Token (JWT) – RFC 7519, https://datatracker.ietf.org/doc/html/rfc7519",
        "[5] OWASP Top Ten, https://owasp.org/www-project-top-ten/",
        "[6] Docker Documentation, https://docs.docker.com/",
        "[7] Jest Testing Framework, https://jestjs.io/",
        "[8] Swagger/OpenAPI Specification, https://swagger.io/specification/",
    ]
    for ref in refs:
        add_paragraph_text(doc, ref, indent=False, font_size=12)

    doc.add_page_break()


# ============================================================
# MAIN
# ============================================================
def main():
    doc = Document()

    # ---- THIẾT LẬP TRANG ----
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    # ---- THIẾT LẬP FONT MẶC ĐỊNH ----
    style = doc.styles['Normal']
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(13)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(6)
    # Đảm bảo font cho tất cả ngôn ngữ
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), "Times New Roman")

    # ---- TẠO NỘI DUNG ----
    create_cover_page(doc)
    create_acknowledgment(doc)
    create_review_pages(doc)
    create_toc_page(doc)
    create_chapter1(doc)
    create_chapter2(doc)
    create_chapter3(doc)
    create_chapter4(doc)
    create_conclusion(doc)
    create_references(doc)

    # ---- SỐ TRANG ----
    add_page_number(doc)

    # ---- LƯU FILE ----
    output_path = "bao_cao_thuc_tap.docx"
    doc.save(output_path)
    print(f"✅ Đã tạo file: {output_path}")
    print("📌 Mở file bằng Word → Nhấn Ctrl+A → F9 để cập nhật mục lục.")


if __name__ == "__main__":
    main()
